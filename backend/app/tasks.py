import os
from celery import Celery
from dotenv import load_dotenv

load_dotenv()
CELERY_BROKER = os.getenv("CELERY_BROKER_URL")
CELERY_BACKEND = os.getenv("CELERY_RESULT_BACKEND")
celery = Celery("worker", broker=CELERY_BROKER, backend=CELERY_BACKEND)

@celery.task
def ingest_document_task(s3_key: str, title: str):
    import asyncio
    asyncio.run(_ingest_document(s3_key, title))

async def _ingest_document(s3_key: str, title: str):
    import boto3
    from .embeddings import embed_text
    from .db import AsyncSessionLocal
    from .models import Document, Chunk
    s3 = boto3.client("s3", endpoint_url=os.getenv("S3_ENDPOINT"),
                      aws_access_key_id=os.getenv("S3_ACCESS_KEY"),
                      aws_secret_access_key=os.getenv("S3_SECRET_KEY"))
    bucket = os.getenv("S3_BUCKET")
    obj = s3.get_object(Bucket=bucket, Key=s3_key)
    raw = obj["Body"].read()
    lower = s3_key.lower()
    text = ""
    if lower.endswith(".pdf"):
        try:
            import fitz
            doc = fitz.open(stream=raw, filetype="pdf")
            pages = [p.get_text() for p in doc]
            text = "\n\n".join(pages)
        except Exception:
            text = ""
    elif lower.endswith(".docx"):
        try:
            import docx
            from io import BytesIO
            d = docx.Document(BytesIO(raw))
            paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs)
        except Exception:
            text = ""
    elif lower.endswith(".xlsx") or lower.endswith(".xls"):
        try:
            import pandas as pd
            from io import BytesIO
            df_map = pd.read_excel(BytesIO(raw), sheet_name=None)
            parts = []
            for sheet, frame in df_map.items():
                parts.append(f"Sheet: {sheet}")
                parts.append(frame.astype(str).fillna("").to_string(index=False))
            text = "\n\n".join(parts)
        except Exception:
            text = ""
    else:
        try:
            text = raw.decode("utf-8")
        except Exception:
            text = raw.decode("latin-1", errors="ignore")

    if not text.strip():
        return

    paragraphs = []
    for p in [p.strip() for p in text.split("\n\n") if p.strip()]:
        if len(p) > 2000:
            for i in range(0, len(p), 1500):
                paragraphs.append(p[i:i+1500])
        else:
            paragraphs.append(p)

    async with AsyncSessionLocal() as session:
        doc = Document(title=title or s3_key, s3_key=s3_key)
        session.add(doc)
        await session.commit()
        await session.refresh(doc)
        for i, para in enumerate(paragraphs):
            emb = await embed_text(para)
            chunk = Chunk(document_id=doc.id, content=para, embedding=emb, metadata={"chunk": i})
            session.add(chunk)
        await session.commit()

@celery.task
def synthesize_answer_task(question: str):
    import asyncio
    return asyncio.run(_synthesize_answer(question))

async def _synthesize_answer(question: str):
    from .embeddings import vector_search, chat_completion
    from .agents import list_agents
    hits = await vector_search(question, top_k=6)
    context = "\n\n".join([f"({h['distance']:.4f}) {h['content']}" for h in hits])
    agents = list_agents()
    outputs = []
    for a in agents:
        out = await chat_completion(a["system"], f"Context:\n{context}\n\nQuestion:\n{question}", temperature=a.get("temperature", 0.1))
        outputs.append((a["name"], out))
    combined = "\n\n".join([f"=={n}==\n{o}" for n, o in outputs])
    synth = await chat_completion("You are a synthesis agent for sales proposals. Produce a concise customer-ready answer and a short Sources section mapping to agent names.", f"Specialist outputs:\n{combined}\n\nQuestion:\n{question}")
    return synth